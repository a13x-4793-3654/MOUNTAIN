"""Excel(.xlsx) テンプレートのプレースホルダ解析と差し込み。

Word と同じ {MOUNTAIN_項目名_連携項目_入力制限_文字制限} 記法を、Excel の
セル（および必要ならヘッダー／フッター）に書いておく。openpyxl でセル値を
走査し、Word 版と共通の PLACEHOLDER_RE / _build_field で解析・置換する。

Excel はセル値が単一の文字列で保持されるため、Word のように 1 つのタグが
複数 run へ分割される問題は起きず、置換はセル単位で完結する。
"""
from __future__ import annotations

import io
from typing import Dict, List

from openpyxl import load_workbook

from .core import PLACEHOLDER_RE, _build_field

# ヘッダー／フッターの左・中央・右パート名。
_HF_ATTRS = ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader", "firstFooter")
_HF_PARTS = ("left", "center", "right")


def _iter_cells(ws):
    for row in ws.iter_rows():
        for cell in row:
            yield cell


def _scan_text(text, seen: Dict[str, dict], order: List[str]) -> None:
    if not isinstance(text, str) or "{MOUNTAIN" not in text:
        return
    for m in PLACEHOLDER_RE.finditer(text):
        token = m.group(0)
        if token in seen:
            continue
        seen[token] = _build_field(m)
        order.append(token)


def parse_template_fields_xlsx(xlsx_bytes: bytes) -> List[dict]:
    """.xlsx を解析し、含まれる差し込み項目（重複除去・出現順）を返す。"""
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    seen: Dict[str, dict] = {}
    order: List[str] = []
    for ws in wb.worksheets:
        for cell in _iter_cells(ws):
            _scan_text(cell.value, seen, order)
        _scan_headers_footers(ws, lambda t: _scan_text(t, seen, order))
    return [seen[t] for t in order]


def _scan_headers_footers(ws, fn) -> None:
    """ヘッダー／フッター各パートの文字列を fn に渡す（openpyxl の版差を吸収）。"""
    try:
        for attr in _HF_ATTRS:
            hf = getattr(ws, attr, None)
            if hf is None:
                continue
            for part_name in _HF_PARTS:
                part = getattr(hf, part_name, None)
                fn(getattr(part, "text", None))
    except Exception:
        # ヘッダー／フッターは補助的なため、解析できなくても本文処理は続行する。
        pass


def _replace_text(text: str, token_to_value: Dict[str, str]) -> str:
    new = text
    for token, value in token_to_value.items():
        if token in new:
            new = new.replace(token, value)
    # 値が渡されなかった残りの MOUNTAIN タグは空文字にして生タグの印字を防ぐ。
    return PLACEHOLDER_RE.sub("", new)


def fill_xlsx(xlsx_bytes: bytes, token_to_value: Dict[str, str]) -> bytes:
    """.xlsx にプレースホルダの値を差し込み、新しい .xlsx バイト列を返す。"""
    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=False)
    for ws in wb.worksheets:
        for cell in _iter_cells(ws):
            v = cell.value
            if isinstance(v, str) and "{MOUNTAIN" in v:
                cell.value = _replace_text(v, token_to_value)
        _fill_headers_footers(ws, token_to_value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _fill_headers_footers(ws, token_to_value: Dict[str, str]) -> None:
    try:
        for attr in _HF_ATTRS:
            hf = getattr(ws, attr, None)
            if hf is None:
                continue
            for part_name in _HF_PARTS:
                part = getattr(hf, part_name, None)
                t = getattr(part, "text", None)
                if isinstance(t, str) and "{MOUNTAIN" in t:
                    part.text = _replace_text(t, token_to_value)
    except Exception:
        pass
